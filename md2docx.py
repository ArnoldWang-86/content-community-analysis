# -*- coding: utf-8 -*-
"""md2docx.py -- 把 Markdown 转成排版良好的 Word 文档
支持: 标题 / 表格 / 代码块 / 有序无序列表 / 复选框 / 引用 / 分隔线 / 粗体 / 行内代码
用法: python md2docx.py <input.md> <output.docx>
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BT = chr(96) # 反引号
FONT = "微软雅黑"
MONO = "Consolas"
ACCENT = "1F3864"
ACCENT2 = "2E5C9A"
CODE_FILL = "F2F3F5"
QUOTE_FILL = "FFF8E6"
INLINE_FILL = "EDEFF2"
TOKEN = re.compile(r"\*\*(.+?)\*\*|" + BT + r"([^" + BT + r"]+)" + BT)


def setfonts(run, ascii_font=FONT, ea_font=FONT):
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:ascii"), ascii_font)
    rf.set(qn("w:hAnsi"), ascii_font)
    rf.set(qn("w:eastAsia"), ea_font)


def shade_p(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_r(run, fill):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    rPr.append(shd)


def border(p, side="bottom", color=ACCENT, sz=8):
    pPr = p._p.get_or_add_pPr()
    b = pPr.find(qn("w:pBdr"))
    if b is None:
        b = OxmlElement("w:pBdr"); pPr.append(b)
    e = OxmlElement("w:" + side)
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "3"); e.set(qn("w:color"), color)
    b.append(e)


def rich(p, text, size=10):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()]); r.font.size = Pt(size)
            setfonts(r)
        if m.group(1) is not None:
            r = p.add_run(m.group(1)); r.font.size = Pt(size); r.bold = True
            setfonts(r)
        else:
            r = p.add_run(m.group(2)); r.font.size = Pt(size - 0.5)
            setfonts(r, MONO, MONO)
            shade_r(r, INLINE_FILL)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:]); r.font.size = Pt(size)
        setfonts(r)


def newdoc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(10)
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:ascii"), FONT); rf.set(qn("w:hAnsi"), FONT); rf.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.15
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.0)
    return doc


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            rich(p, txt, 9)
            if ri == 0:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor.from_string("FFFFFF")
                tcPr = cells[ci]._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), ACCENT)
                tcPr.append(shd)
    return t


def convert(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = newdoc()
    i, n = 0, len(lines)
    fence = BT * 3

    while i < n:
        raw = lines[i]
        st = raw.strip()

        if st.startswith(fence):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith(fence):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(4); pf.space_after = Pt(7)
            pf.left_indent = Cm(0.3); pf.right_indent = Cm(0.1)
            pf.line_spacing = 1.0
            shade_p(p, CODE_FILL)
            border(p, "left", "9AA5B1", 18)
            for k, cl in enumerate(buf):
                if k:
                    p.add_run().add_break()
                r = p.add_run(cl)
                r.font.size = Pt(8.5)
                setfonts(r, MONO, MONO)
                r.font.color.rgb = RGBColor.from_string("1F2328")
            continue

        if st.startswith("|") and st.endswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            rows = []
            for bi, bl in enumerate(block):
                cells = [c.strip() for c in bl.strip("|").split("|")]
                if bi == 1 and all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                add_table(doc, rows)
                sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
            continue

        if re.fullmatch(r"-{3,}", st):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
            border(p, "bottom", "C6CBD1", 6)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", st)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            sizes = {1: 16, 2: 13.5, 3: 11.5, 4: 10.5}
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(14 if lvl <= 2 else 9)
            pf.space_after = Pt(5)
            rich(p, txt, sizes.get(lvl, 10.5))
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(ACCENT if lvl <= 2 else ACCENT2)
            if lvl <= 2:
                border(p, "bottom", ACCENT if lvl == 1 else "B8C4D9", 8 if lvl == 1 else 6)
            i += 1
            continue

        if st.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(4); pf.space_after = Pt(7)
            pf.left_indent = Cm(0.35); pf.right_indent = Cm(0.1)
            shade_p(p, QUOTE_FILL)
            border(p, "left", "E0A32E", 18)
            first = True
            for bl in buf:
                if not first:
                    p.add_run().add_break()
                first = False
                rich(p, bl.strip(), 9.5)
            continue

        mb = re.match(r"^(\s*)[-*]\s+(.*)$", raw)
        mo = re.match(r"^(\s*)(\d+)\.\s+(.*)$", raw)
        if mb or mo:
            if mb:
                indent = len(mb.group(1)) // 2
                body = mb.group(2)
                cb = re.match(r"^\[( |x|X)\]\s*(.*)$", body)
                if cb:
                    marker = "☑ " if cb.group(1).lower() == "x" else "☐ "
                    body = cb.group(2)
                else:
                    marker = "• "
            else:
                indent = len(mo.group(1)) // 2
                marker = mo.group(2) + ". "
                body = mo.group(3)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.55 + indent * 0.5)
            pf.first_line_indent = Cm(-0.55)
            pf.space_before = Pt(1); pf.space_after = Pt(2)
            r = p.add_run(marker); r.font.size = Pt(10)
            setfonts(r)
            rich(p, body, 10)
            i += 1
            continue

        if not st:
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        rich(p, st, 10)
        i += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    print("OK", convert(sys.argv[1], sys.argv[2]))
